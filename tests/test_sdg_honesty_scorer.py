import http.client
import io
import json
import tempfile
import unittest
from pathlib import Path
from typing import Any, cast
from urllib.request import Request

from sdg_honesty_scorer import (
    DEFAULT_OPENROUTER_MODEL,
    AIClient,
    InvestmentCategorization,
    KeywordCategoryItem,
    ReportInput,
    SdgMentionItem,
    SdgMentions,
    _merge_mentions,
    _safe_excel_value,
    compute_mismatch,
    create_excel,
    discover_pdfs,
    load_report_manifest,
    process_bank_reports,
    process_reports,
)


class FakeResponse(io.BytesIO):
    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, tb):
        return False


class FakeOpen:
    def __init__(self, content: str):
        self.content = content
        self.requests: list[Request] = []

    def __call__(self, request: Request, timeout: float):
        self.requests.append(request)
        body = {
            "choices": [
                {
                    "message": {
                        "content": self.content,
                    }
                }
            ]
        }
        return FakeResponse(json.dumps(body).encode("utf-8"))


class TimeoutThenSuccessOpen(FakeOpen):
    def __call__(self, request: Request, timeout: float):
        if not self.requests:
            self.requests.append(request)
            raise TimeoutError("timed out")
        return super().__call__(request, timeout)


class ConnectionResetThenSuccessOpen(FakeOpen):
    def __call__(self, request: Request, timeout: float):
        if not self.requests:
            self.requests.append(request)
            raise ConnectionResetError("connection reset by peer")
        return super().__call__(request, timeout)


class NonJsonResponseThenSuccessOpen(FakeOpen):
    def __call__(self, request: Request, timeout: float):
        if not self.requests:
            self.requests.append(request)
            return FakeResponse(b"<html>temporary upstream failure</html>")
        return super().__call__(request, timeout)


class IncompleteReadThenSuccessOpen(FakeOpen):
    def __call__(self, request: Request, timeout: float):
        if not self.requests:
            self.requests.append(request)
            raise http.client.IncompleteRead(b"partial")
        return super().__call__(request, timeout)


class SequenceOpen:
    def __init__(self, contents: list[object]):
        self.contents = contents
        self.requests: list[Request] = []

    def __call__(self, request: Request, timeout: float):
        self.requests.append(request)
        index = min(len(self.requests) - 1, len(self.contents) - 1)
        body = {"choices": [{"message": {"content": self.contents[index]}}]}
        return FakeResponse(json.dumps(body).encode("utf-8"))


class OpenRouterClientTests(unittest.TestCase):
    def test_fetch_sdg_mentions_uses_openrouter_and_normalizes_json(self):
        fake_open = FakeOpen(
            "```json\n"
            '{"sdg_data":[{"sdg_number":7,"mention_count":2,'
            '"importance_score":4.5,"page_numbers":[3,"4",3]}]}'
            "\n```"
        )
        client = AIClient(api_key="test-key", opener=fake_open)

        mentions = client.fetch_sdg_mentions("[PAGE 3]\nRenewable energy finance")

        self.assertEqual(
            mentions,
            [
                {
                    "sdg_number": 7,
                    "mention_count": 2,
                    "importance_score": 4.5,
                    "page_numbers": [3, 4],
                }
            ],
        )
        self.assertEqual(len(fake_open.requests), 1)
        request = fake_open.requests[0]
        self.assertEqual(request.full_url, "https://openrouter.ai/api/v1/chat/completions")
        self.assertEqual(request.get_header("Authorization"), "Bearer test-key")
        payload = json.loads(request.data.decode("utf-8"))
        self.assertEqual(payload["model"], DEFAULT_OPENROUTER_MODEL)
        self.assertEqual(payload["max_tokens"], 8192)
        self.assertEqual(payload["response_format"], {"type": "json_object"})
        self.assertEqual(payload["messages"][-1]["role"], "user")

    def test_categorize_investments_normalizes_missing_fields(self):
        fake_open = FakeOpen(
            json.dumps(
                {
                    "keyword_categories": [
                        {
                            "category": "Renewable energy",
                            "amount_eur": "1250000",
                            "relevant_sdgs": [7, "13", 99],
                            "page_numbers": ["3", 3],
                            "evidence": "Portfolio finances renewable capacity.",
                        }
                    ],
                    "investment_instruments": [
                        {
                            "instrument": "Green bond",
                            "subcategories": "solar",
                            "amount_eur": 500000,
                            "relevant_sdgs": [7],
                            "justification": "Funds solar projects.",
                            "page_numbers": [4],
                            "evidence": "Green bond proceeds funded solar.",
                        }
                    ],
                    "unlinked_investments_eur": "1000.50",
                }
            )
        )
        client = AIClient(api_key="test-key", opener=fake_open)

        categorized = client.categorize_investments("portfolio text")

        category = categorized["keyword_categories"][0]
        instrument = categorized["investment_instruments"][0]
        self.assertEqual(category["category"], "Renewable energy")
        self.assertEqual(category["subcategories"], [])
        self.assertEqual(category["amount_eur"], 1250000.0)
        self.assertEqual(category["relevant_sdgs"], [7, 13])
        self.assertEqual(category["justification"], "")
        self.assertEqual(category["page_numbers"], [3])
        self.assertEqual(category["evidence"], "Portfolio finances renewable capacity.")
        self.assertEqual(category["construct_type"], "unknown")
        self.assertEqual(category["sdg_support"], "none")
        self.assertEqual(category["inclusion_policy"], "excluded")
        self.assertFalse(category["score_eligible"])
        self.assertIn("unknown_construct", category["exclusion_reason"])
        self.assertEqual(instrument["instrument"], "Green bond")
        self.assertEqual(instrument["subcategories"], ["solar"])
        self.assertEqual(instrument["inclusion_policy"], "excluded")
        self.assertEqual(categorized["unlinked_investments_eur"], 1000.50)

    def test_categorize_investments_converts_eur_thousand_source_amount_for_scoring(self):
        fake_open = FakeOpen(
            json.dumps(
                {
                    "keyword_categories": [
                        {
                            "category": "Reconstruction finance",
                            "amount_eur": 300007000000,
                            "source_amount_text": "EUR 300,000 thousand",
                            "source_value": "300000",
                            "source_unit": "thousand",
                            "source_currency": "EUR",
                            "construct_type": "disclosed_sdg_linked_finance",
                            "sdg_support": "explicit",
                            "relevant_sdgs": [8],
                            "page_numbers": [5],
                            "evidence": "The financing agreement totals EUR 300,000 thousand.",
                        }
                    ],
                    "investment_instruments": [],
                }
            )
        )
        client = AIClient(api_key="test-key", opener=fake_open)

        row = client.categorize_investments("[PAGE 5]\nThe financing agreement totals EUR 300,000 thousand.")[
            "keyword_categories"
        ][0]

        self.assertEqual(row["amount_eur"], 300000000.0)
        self.assertEqual(row["source_currency"], "EUR")
        self.assertEqual(row["source_unit"], "thousand")
        self.assertEqual(row["inclusion_policy"], "linked_scored")
        self.assertTrue(row["score_eligible"])
        self.assertEqual(row["exclusion_reason"], "")

    def test_categorize_investments_converts_non_eur_with_review_flag(self):
        fake_open = FakeOpen(
            json.dumps(
                {
                    "investment_instruments": [
                        {
                            "instrument": "Sustainable financing",
                            "amount_eur": 361539000000,
                            "source_amount_text": "DKK 365bn in sustainable financing",
                            "source_value": 365,
                            "source_unit": "billion",
                            "source_currency": "DKK",
                            "construct_type": "disclosed_sdg_linked_finance",
                            "sdg_support": "explicit",
                            "relevant_sdgs": [7, 13],
                            "page_numbers": [13, 14],
                            "evidence": "DKK 365bn in sustainable financing, including green loans and arranged bonds.",
                        }
                    ]
                }
            )
        )
        client = AIClient(api_key="test-key", opener=fake_open)

        row = client.categorize_investments("[PAGE 13]\nDKK 365bn in sustainable financing")["investment_instruments"][
            0
        ]

        self.assertEqual(row["source_currency"], "DKK")
        self.assertEqual(row["amount_eur"], 48910000000.0)
        self.assertNotEqual(row["amount_eur"], 361539000000.0)
        self.assertTrue(row["needs_fx_review"])
        self.assertEqual(row["conversion_rate_used"], 0.134)
        self.assertEqual(row["conversion_source"], "fixed_fx_table_2024-12-31")

    def test_categorize_investments_uses_report_year_fx_rate_when_known(self):
        fake_open = FakeOpen(
            json.dumps(
                {
                    "investment_instruments": [
                        {
                            "instrument": "Sustainable financing",
                            "source_amount_text": "DKK 365bn in sustainable financing",
                            "source_value": 365,
                            "source_unit": "billion",
                            "source_currency": "DKK",
                            "construct_type": "disclosed_sdg_linked_finance",
                            "sdg_support": "explicit",
                            "relevant_sdgs": [7, 13],
                            "page_numbers": [13, 14],
                            "evidence": "DKK 365bn in sustainable financing, including green loans and arranged bonds.",
                        }
                    ]
                }
            )
        )
        client = AIClient(api_key="test-key", opener=fake_open)

        row = client.categorize_investments("[PAGE 13]\nDKK 365bn in sustainable financing", "2021")[
            "investment_instruments"
        ][0]

        self.assertEqual(row["source_currency"], "DKK")
        # 2021 ECB year-end DKK rate (0.13447) differs from the fixed fallback (0.134).
        self.assertEqual(row["conversion_rate_used"], 0.13447)
        self.assertEqual(row["conversion_source"], "ecb_nbu_year_end_reference_rate:2021")
        self.assertAlmostEqual(row["amount_eur"], 365 * 1_000_000_000 * 0.13447, delta=1.0)
        self.assertTrue(row["needs_fx_review"])

    def test_categorize_investments_falls_back_to_fixed_rate_for_unmapped_year(self):
        fake_open = FakeOpen(
            json.dumps(
                {
                    "investment_instruments": [
                        {
                            "instrument": "Sustainable financing",
                            "source_amount_text": "DKK 365bn in sustainable financing",
                            "source_value": 365,
                            "source_unit": "billion",
                            "source_currency": "DKK",
                            "construct_type": "disclosed_sdg_linked_finance",
                            "sdg_support": "explicit",
                            "relevant_sdgs": [7, 13],
                            "page_numbers": [13, 14],
                            "evidence": "DKK 365bn in sustainable financing, including green loans and arranged bonds.",
                        }
                    ]
                }
            )
        )
        client = AIClient(api_key="test-key", opener=fake_open)

        row = client.categorize_investments("[PAGE 13]\nDKK 365bn in sustainable financing", "1999")[
            "investment_instruments"
        ][0]

        self.assertEqual(row["conversion_rate_used"], 0.134)
        self.assertEqual(row["conversion_source"], "fixed_fx_table_2024-12-31:fallback_unmapped_year_1999")

    def test_categorize_investments_downgrades_facilitated_on_behalf_of_customers(self):
        fake_open = FakeOpen(
            json.dumps(
                {
                    "investment_instruments": [
                        {
                            "instrument": "Green bonds",
                            "amount_eur": 25728000000.0,
                            "source_amount_text": "DKK 192 billion",
                            "source_value": 192,
                            "source_unit": "billion",
                            "source_currency": "DKK",
                            "construct_type": "facilitated_on_behalf_of_customers",
                            "sdg_support": "strong",
                            "relevant_sdgs": [7, 13],
                            "page_numbers": [10],
                            "evidence": "DKK 192 billion in green bonds issued on behalf of customers.",
                        }
                    ]
                }
            )
        )
        client = AIClient(api_key="test-key", opener=fake_open)

        row = client.categorize_investments("[PAGE 10]\nDKK 192 billion issued on behalf of customers")[
            "investment_instruments"
        ][0]

        self.assertEqual(row["inclusion_policy"], "gross_context_only")
        self.assertFalse(row["score_eligible"])
        self.assertIn("context_construct", row["exclusion_reason"])

    def test_categorize_investments_excludes_forward_commitment(self):
        fake_open = FakeOpen(
            json.dumps(
                {
                    "keyword_categories": [
                        {
                            "category": "Green loans",
                            "amount_eur": 400000000.0,
                            "source_amount_text": "EUR 400m",
                            "source_value": 400,
                            "source_unit": "million",
                            "source_currency": "EUR",
                            "construct_type": "forward_commitment",
                            "sdg_support": "strong",
                            "relevant_sdgs": [7],
                            "page_numbers": [8],
                            "evidence": "We have committed to issuing up to EUR 400m within the next three years.",
                        }
                    ]
                }
            )
        )
        client = AIClient(api_key="test-key", opener=fake_open)

        row = client.categorize_investments("[PAGE 8]\ncommitted to issuing up to EUR 400m")["keyword_categories"][0]

        self.assertEqual(row["inclusion_policy"], "excluded")
        self.assertFalse(row["score_eligible"])
        self.assertIn("excluded_construct", row["exclusion_reason"])

    def test_categorize_investments_excludes_recipient_side_financing(self):
        fake_open = FakeOpen(
            json.dumps(
                {
                    "investment_instruments": [
                        {
                            "instrument": "IBRD loan",
                            "amount_eur": 120794088.0,
                            "source_amount_text": "UAH 5,033,087 thousand",
                            "source_value": 5033087,
                            "source_unit": "thousand",
                            "source_currency": "UAH",
                            "construct_type": "recipient_side_financing",
                            "sdg_support": "strong",
                            "relevant_sdgs": [7, 13],
                            "page_numbers": [97],
                            "evidence": (
                                "Loans from international financial institutions include the loan from the IBRD "
                                "under the Project on Energy Efficiency."
                            ),
                        }
                    ]
                }
            )
        )
        client = AIClient(api_key="test-key", opener=fake_open)

        row = client.categorize_investments("[PAGE 97]\nloan from the IBRD")["investment_instruments"][0]

        self.assertEqual(row["inclusion_policy"], "excluded")
        self.assertFalse(row["score_eligible"])
        self.assertIn("excluded_construct:recipient_side_financing", row["exclusion_reason"])

    def test_categorize_investments_downgrades_bare_credit_limit_without_financing_verb(self):
        fake_open = FakeOpen(
            json.dumps(
                {
                    "keyword_categories": [
                        {
                            "category": "Guarantees",
                            "amount_eur": 75000000.0,
                            "source_amount_text": "up to 75 million euros",
                            "source_value": 75,
                            "source_unit": "million",
                            "source_currency": "EUR",
                            "construct_type": "credit_limit_undrawn",
                            "sdg_support": "strong",
                            "relevant_sdgs": [8],
                            "page_numbers": [4],
                            "evidence": "The Bank offers a non-financial guarantee limit of up to 75 million euros.",
                        }
                    ]
                }
            )
        )
        client = AIClient(api_key="test-key", opener=fake_open)

        row = client.categorize_investments("[PAGE 4]\nguarantee limit of up to 75 million euros")[
            "keyword_categories"
        ][0]

        self.assertEqual(row["inclusion_policy"], "gross_context_only")
        self.assertFalse(row["score_eligible"])

    def test_categorize_investments_financing_verb_does_not_override_credit_limit(self):
        fake_open = FakeOpen(
            json.dumps(
                {
                    "keyword_categories": [
                        {
                            "category": "Sustainable business loans",
                            "amount_eur": 72000000.0,
                            "source_amount_text": "over UAH 3 billion",
                            "source_value": 3,
                            "source_unit": "billion",
                            "source_currency": "UAH",
                            "construct_type": "credit_limit_undrawn",
                            "sdg_support": "explicit",
                            "relevant_sdgs": [7, 13],
                            "page_numbers": [132],
                            "evidence": (
                                "In 2023, 229 credit agreements for sustainable business development projects "
                                "were financed for a total amount (credit limits) of over UAH 3 billion."
                            ),
                        }
                    ]
                }
            )
        )
        client = AIClient(api_key="test-key", opener=fake_open)

        row = client.categorize_investments("[PAGE 132]\n229 credit agreements were financed")["keyword_categories"][0]

        self.assertEqual(row["inclusion_policy"], "gross_context_only")
        self.assertFalse(row["score_eligible"])
        self.assertEqual(row["exclusion_reason"], "context_construct:credit_limit_undrawn")

    def test_categorize_investments_keeps_facilitated_finance_as_gross_context_only(self):
        fake_open = FakeOpen(
            json.dumps(
                {
                    "keyword_categories": [
                        {
                            "category": "Sustainable finance aggregate",
                            "amount_eur": 1429000000000,
                            "source_amount_text": "facilitated EUR 135bn in sustainable financing over 2022-2023",
                            "source_value": 135,
                            "source_unit": "billion",
                            "source_currency": "EUR",
                            "construct_type": "facilitated_sustainable_finance",
                            "sdg_support": "strong",
                            "relevant_sdgs": [13],
                            "page_numbers": [6, 9],
                            "evidence": "Nordea facilitated EUR 135bn in sustainable financing over 2022-2023.",
                        }
                    ]
                }
            )
        )
        client = AIClient(api_key="test-key", opener=fake_open)

        row = client.categorize_investments("[PAGE 6]\nfacilitated EUR 135bn in sustainable financing")[
            "keyword_categories"
        ][0]

        self.assertEqual(row["amount_eur"], 135000000000.0)
        self.assertEqual(row["inclusion_policy"], "gross_context_only")
        self.assertFalse(row["score_eligible"])
        self.assertIn("context_construct", row["exclusion_reason"])

    def test_categorize_investments_normalizes_euro_spelled_out_currency(self):
        fake_open = FakeOpen(
            json.dumps(
                {
                    "keyword_categories": [
                        {
                            "category": "Renewable energy",
                            "source_amount_text": "83.1 mln. Euro",
                            "source_value": 83.1,
                            "source_unit": "million",
                            "source_currency": "Euro",
                            "construct_type": "facilitated_sustainable_finance",
                            "sdg_support": "strong",
                            "relevant_sdgs": [7, 13],
                            "page_numbers": [5],
                            "evidence": "limits for financing of 30 projects ... for 83.1 mln. Euro",
                        }
                    ]
                }
            )
        )
        client = AIClient(api_key="test-key", opener=fake_open)

        row = client.categorize_investments("[PAGE 5]\n83.1 mln. Euro")["keyword_categories"][0]

        self.assertEqual(row["source_currency"], "EUR")
        self.assertEqual(row["amount_eur"], 83100000.0)
        self.assertFalse(row["needs_fx_review"])

    def test_categorize_investments_normalizes_other_spelled_out_currencies(self):
        cases = [
            ("US Dollars", "USD"),
            ("Pounds Sterling", "GBP"),
            ("Norwegian Kroner", "NOK"),
            ("Swedish Krona", "SEK"),
            ("Danish Kroner", "DKK"),
            ("Ukrainian Hryvnia", "UAH"),
        ]
        for spelled_out, expected_code in cases:
            with self.subTest(spelled_out=spelled_out):
                fake_open = FakeOpen(
                    json.dumps(
                        {
                            "keyword_categories": [
                                {
                                    "category": "Renewable energy",
                                    "source_amount_text": f"10 million {spelled_out}",
                                    "source_value": 10,
                                    "source_unit": "million",
                                    "source_currency": spelled_out,
                                    "construct_type": "disclosed_sdg_linked_finance",
                                    "sdg_support": "explicit",
                                    "relevant_sdgs": [7],
                                    "page_numbers": [1],
                                    "evidence": f"10 million {spelled_out} for renewable energy.",
                                }
                            ]
                        }
                    )
                )
                client = AIClient(api_key="test-key", opener=fake_open)

                row = client.categorize_investments(f"[PAGE 1]\n10 million {spelled_out}")["keyword_categories"][0]

                self.assertEqual(row["source_currency"], expected_code)

    def test_categorize_investments_excludes_generic_portfolio_growth(self):
        fake_open = FakeOpen(
            json.dumps(
                {
                    "keyword_categories": [
                        {
                            "category": "Corporate loan book",
                            "source_amount_text": "portfolio increased by UAH 15.9bn",
                            "source_value": 15.9,
                            "source_unit": "billion",
                            "source_currency": "UAH",
                            "construct_type": "portfolio_share_or_growth",
                            "sdg_support": "none",
                            "relevant_sdgs": [8],
                            "page_numbers": [2],
                            "evidence": (
                                "The bank reported corporate-lending market share and portfolio growth of UAH 15.9bn."
                            ),
                        }
                    ]
                }
            )
        )
        client = AIClient(api_key="test-key", opener=fake_open)

        row = client.categorize_investments("[PAGE 2]\nportfolio increased by UAH 15.9bn")["keyword_categories"][0]

        self.assertEqual(row["inclusion_policy"], "excluded")
        self.assertFalse(row["score_eligible"])
        self.assertIn("excluded_construct", row["exclusion_reason"])

    def test_categorize_investments_retries_timeout_error(self):
        fake_open = TimeoutThenSuccessOpen(
            json.dumps(
                {
                    "keyword_categories": [],
                    "investment_instruments": [],
                    "unlinked_investments_eur": 0,
                }
            )
        )
        client = AIClient(api_key="test-key", opener=fake_open, max_retries=1, sleep=lambda _: None)

        categorized = client.categorize_investments("portfolio text")

        self.assertEqual(categorized["keyword_categories"], [])
        self.assertEqual(len(fake_open.requests), 2)

    def test_categorize_investments_retries_connection_reset_error(self):
        fake_open = ConnectionResetThenSuccessOpen(
            json.dumps(
                {
                    "keyword_categories": [],
                    "investment_instruments": [],
                    "unlinked_investments_eur": 0,
                }
            )
        )
        client = AIClient(api_key="test-key", opener=fake_open, max_retries=1, sleep=lambda _: None)

        categorized = client.categorize_investments("portfolio text")

        self.assertEqual(categorized["keyword_categories"], [])
        self.assertEqual(len(fake_open.requests), 2)

    def test_categorize_investments_retries_non_json_openrouter_response(self):
        fake_open = NonJsonResponseThenSuccessOpen(
            json.dumps(
                {
                    "keyword_categories": [],
                    "investment_instruments": [],
                    "unlinked_investments_eur": 0,
                }
            )
        )
        client = AIClient(api_key="test-key", opener=fake_open, max_retries=1, sleep=lambda _: None)

        categorized = client.categorize_investments("portfolio text")

        self.assertEqual(categorized["keyword_categories"], [])
        self.assertEqual(len(fake_open.requests), 2)

    def test_categorize_investments_retries_incomplete_read_error(self):
        fake_open = IncompleteReadThenSuccessOpen(
            json.dumps(
                {
                    "keyword_categories": [],
                    "investment_instruments": [],
                    "unlinked_investments_eur": 0,
                }
            )
        )
        client = AIClient(api_key="test-key", opener=fake_open, max_retries=1, sleep=lambda _: None)

        categorized = client.categorize_investments("portfolio text")

        self.assertEqual(categorized["keyword_categories"], [])
        self.assertEqual(len(fake_open.requests), 2)

    def test_categorize_investments_retries_invalid_model_json(self):
        fake_open = SequenceOpen(
            [
                '{"keyword_categories"',
                json.dumps(
                    {
                        "keyword_categories": [],
                        "investment_instruments": [],
                        "unlinked_investments_eur": 0,
                    }
                ),
            ]
        )
        client = AIClient(api_key="test-key", opener=fake_open, max_retries=1, sleep=lambda _: None)

        categorized = client.categorize_investments("portfolio text")

        self.assertEqual(categorized["keyword_categories"], [])
        self.assertEqual(len(fake_open.requests), 2)

    def test_categorize_investments_retries_null_message_content(self):
        fake_open = SequenceOpen(
            [
                None,
                json.dumps(
                    {
                        "keyword_categories": [],
                        "investment_instruments": [],
                        "unlinked_investments_eur": 0,
                    }
                ),
            ]
        )
        client = AIClient(api_key="test-key", opener=fake_open, max_retries=1, sleep=lambda _: None)

        categorized = client.categorize_investments("portfolio text")

        self.assertEqual(categorized["keyword_categories"], [])
        self.assertEqual(len(fake_open.requests), 2)

    def test_rejects_non_openrouter_base_url_by_default(self):
        with self.assertRaisesRegex(RuntimeError, "OPENROUTER_BASE_URL"):
            AIClient(api_key="test-key", base_url="http://example.com/chat")

    def test_rejects_proprietary_models_on_openrouter(self):
        blocked = [
            ("openai/gpt-4o", "direct OpenAI"),
            ("openai/o3", "direct OpenAI"),
            ("anthropic/claude-sonnet-5-5", "direct Anthropic"),
            ("google/gemini-2.5-pro", "direct Google"),
        ]
        for model, message in blocked:
            with self.subTest(model=model), self.assertRaisesRegex(RuntimeError, message):
                AIClient(model=model, api_key="test-key")

    def test_requires_openrouter_api_key(self):
        with self.assertRaisesRegex(RuntimeError, "OPENROUTER_API_KEY"):
            AIClient(api_key="")


class OutputSafetyTests(unittest.TestCase):
    def test_safe_excel_value_escapes_formula_like_strings(self):
        self.assertEqual(_safe_excel_value('=HYPERLINK("http://evil")'), '\'=HYPERLINK("http://evil")')
        self.assertEqual(_safe_excel_value("  +SUM(1,1)"), "'  +SUM(1,1)")
        self.assertEqual(_safe_excel_value("normal text"), "normal text")
        self.assertEqual(_safe_excel_value(123), 123)

    def test_long_bank_names_preserve_distinct_excel_sheet_suffixes(self):
        bank = "Finland_Säästöpankkien Keskuspankki_2024"
        mentions = {
            bank: SdgMentions([SdgMentionItem(sdg_number=7, mention_count=1, importance_score=2.0, page_numbers=[1])])
        }
        investments = {
            bank: InvestmentCategorization(
                keyword_categories=[
                    KeywordCategoryItem(
                        category="Renewable energy",
                        subcategories=[],
                        amount_eur=0.0,
                        relevant_sdgs=[],
                    )
                ]
            )
        }
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "long_names.xlsx"
            create_excel(mentions, investments, str(output))
            import openpyxl

            workbook = openpyxl.load_workbook(output, read_only=True)
            self.assertIn("Finland_Säästöpankkien_Mentions", workbook.sheetnames)
            self.assertIn("Finland_Säästöpankkien_Keywords", workbook.sheetnames)
            self.assertIn("Finland_Säästöpankk_Instruments", workbook.sheetnames)


class ProcessingFailureTests(unittest.TestCase):
    def test_model_batch_failure_fails_the_run(self):
        class FailingClient:
            def fetch_sdg_mentions(self, text):
                raise RuntimeError("model unavailable")

            def categorize_investments(self, text, report_year=""):
                return {
                    "keyword_categories": [],
                    "investment_instruments": [],
                    "unlinked_investments_eur": 0.0,
                }

        with tempfile.TemporaryDirectory() as tmp:
            pdf_path = f"{tmp}/Bank_report.pdf"
            cache_path = f"{pdf_path}.parsed.json"
            with open(cache_path, "w", encoding="utf-8") as file:
                json.dump({"full_text": "report text", "pages": [[1, "report text"]]}, file)

            with self.assertRaisesRegex(RuntimeError, "SDG extraction failed"):
                process_bank_reports(FailingClient(), [pdf_path])

    def test_investment_categorization_receives_page_labels(self):
        class RecordingClient:
            def __init__(self):
                self.investment_texts = []

            def fetch_sdg_mentions(self, text):
                return []

            def categorize_investments(self, text, report_year=""):
                self.investment_texts.append(text)
                return {
                    "keyword_categories": [],
                    "investment_instruments": [],
                    "unlinked_investments_eur": 0.0,
                }

        with tempfile.TemporaryDirectory() as tmp:
            pdf_path = f"{tmp}/Bank_report.pdf"
            cache_path = f"{pdf_path}.parsed.json"
            with open(cache_path, "w", encoding="utf-8") as file:
                json.dump({"full_text": "first\nsecond", "pages": [[1, "first"], [2, "second"]]}, file)

            client = RecordingClient()
            process_bank_reports(client, [pdf_path])

            investment_text = "\n".join(client.investment_texts)
            self.assertIn("[PAGE 1]\nfirst", investment_text)
            self.assertIn("[PAGE 2]\nsecond", investment_text)

    def test_investment_categorization_is_chunked_and_merged(self):
        class RecordingClient:
            def __init__(self):
                self.investment_texts = []

            def fetch_sdg_mentions(self, text):
                return []

            def categorize_investments(self, text, report_year=""):
                self.investment_texts.append(text)
                return {
                    "keyword_categories": [
                        {
                            "category": "Renewable energy",
                            "subcategories": [],
                            "amount_eur": 1000000.0,
                            "source_amount_text": "EUR 1 million",
                            "source_value": 1,
                            "source_unit": "million",
                            "source_currency": "EUR",
                            "construct_type": "disclosed_sdg_linked_finance",
                            "sdg_support": "explicit",
                            "relevant_sdgs": [7],
                            "page_numbers": [1],
                            "evidence": "EUR 1 million for renewable energy.",
                        }
                    ],
                    "investment_instruments": [],
                    "unlinked_investments_eur": 2.0,
                }

        with tempfile.TemporaryDirectory() as tmp:
            pdf_path = f"{tmp}/Bank_report.pdf"
            cache_path = f"{pdf_path}.parsed.json"
            pages = [[1, "a" * 16000], [2, "b" * 16000]]
            with open(cache_path, "w", encoding="utf-8") as file:
                json.dump({"full_text": "", "pages": pages}, file)

            client = RecordingClient()
            _mentions, investments = process_bank_reports(client, [pdf_path])

            self.assertEqual(len(client.investment_texts), 2)
            self.assertEqual(len(investments["Bank"].keyword_categories), 2)
            self.assertEqual(investments["Bank"].unlinked_investments_eur, 4.0)

    def test_process_reports_accepts_max_workers_for_chunked_investments(self):
        class RecordingClient:
            def __init__(self):
                self.investment_texts = []

            def fetch_sdg_mentions(self, text):
                return []

            def categorize_investments(self, text, report_year=""):
                self.investment_texts.append(text)
                return {"keyword_categories": [], "investment_instruments": [], "unlinked_investments_eur": 0.0}

        with tempfile.TemporaryDirectory() as tmp:
            pdf_path = f"{tmp}/Bank_report.pdf"
            cache_path = f"{pdf_path}.parsed.json"
            pages = [[1, "a" * 16000], [2, "b" * 16000]]
            with open(cache_path, "w", encoding="utf-8") as file:
                json.dump({"full_text": "", "pages": pages}, file)

            client = RecordingClient()
            process_bank_reports(client, [pdf_path], max_workers=2)

            self.assertEqual(len(client.investment_texts), 2)

    def test_process_reports_logs_bank_and_chunk_progress(self):
        class RecordingClient:
            def fetch_sdg_mentions(self, text):
                return []

            def categorize_investments(self, text, report_year=""):
                return {"keyword_categories": [], "investment_instruments": [], "unlinked_investments_eur": 0.0}

        with tempfile.TemporaryDirectory() as tmp:
            pdf_path = f"{tmp}/Bank_report.pdf"
            cache_path = f"{pdf_path}.parsed.json"
            with open(cache_path, "w", encoding="utf-8") as file:
                json.dump({"full_text": "report text", "pages": [[1, "report text"]]}, file)

            with self.assertLogs("honesty_scorer.pipeline", level="INFO") as logs:
                process_bank_reports(RecordingClient(), [pdf_path])

            output = "\n".join(logs.output)
            self.assertIn("bank_start bank=Bank", output)
            self.assertIn("sdg_chunk_start bank=Bank chunk=1/1", output)
            self.assertIn("investment_chunk_start bank=Bank chunk=1/1", output)
            self.assertIn("bank_complete bank=Bank", output)

    def test_process_reports_checkpoints_successful_bank_before_later_failure(self):
        class FailingSecondBankClient:
            def fetch_sdg_mentions(self, text):
                if "beta report" in text:
                    raise RuntimeError("beta failed")
                return [{"sdg_number": 7, "mention_count": 1, "importance_score": 4.0, "page_numbers": [1]}]

            def categorize_investments(self, text, report_year=""):
                return {"keyword_categories": [], "investment_instruments": [], "unlinked_investments_eur": 0.0}

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            alpha = root / "Alpha.pdf"
            beta = root / "Beta.pdf"
            (root / "Alpha.pdf.parsed.json").write_text(
                json.dumps({"full_text": "alpha report", "pages": [[1, "alpha report"]]}), encoding="utf-8"
            )
            (root / "Beta.pdf.parsed.json").write_text(
                json.dumps({"full_text": "beta report", "pages": [[1, "beta report"]]}), encoding="utf-8"
            )
            checkpoint_dir = root / "checkpoints"
            reports = [
                ReportInput(str(alpha), bank="Alpha", country="finland", year="2024"),
                ReportInput(str(beta), bank="Beta", country="finland", year="2024"),
            ]

            with self.assertRaisesRegex(RuntimeError, "SDG extraction failed for finland_Beta_2024"):
                process_reports(FailingSecondBankClient(), reports, checkpoint_dir=checkpoint_dir)

            checkpoints = sorted(
                path for path in checkpoint_dir.glob("*.json") if not path.name.endswith(".error.json")
            )
            errors = sorted(checkpoint_dir.glob("*.error.json"))
            self.assertEqual(len(checkpoints), 1)
            self.assertEqual(len(errors), 1)
            saved = json.loads(checkpoints[0].read_text(encoding="utf-8"))
            self.assertEqual(saved["bank"], "finland_Alpha_2024")
            sdg_7 = [item for item in saved["mentions"]["sdg_data"] if item["sdg_number"] == 7][0]
            self.assertEqual(sdg_7["mention_count"], 1)

    def test_successful_checkpoint_removes_stale_error_marker(self):
        class FailingThenSuccessfulClient:
            def __init__(self, should_fail: bool):
                self.should_fail = should_fail

            def fetch_sdg_mentions(self, text):
                if self.should_fail:
                    raise RuntimeError("temporary model failure")
                return [{"sdg_number": 7, "mention_count": 1, "importance_score": 4.0, "page_numbers": [1]}]

            def categorize_investments(self, text, report_year=""):
                return {"keyword_categories": [], "investment_instruments": [], "unlinked_investments_eur": 0.0}

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            report = root / "Bank.pdf"
            (root / "Bank.pdf.parsed.json").write_text(
                json.dumps({"full_text": "bank report", "pages": [[1, "bank report"]]}), encoding="utf-8"
            )
            checkpoint_dir = root / "checkpoints"
            report_input = ReportInput(str(report), bank="Bank", country="finland", year="2024")

            with self.assertRaisesRegex(RuntimeError, "SDG extraction failed"):
                process_reports(
                    cast(Any, FailingThenSuccessfulClient(should_fail=True)),
                    [report_input],
                    checkpoint_dir=checkpoint_dir,
                )
            self.assertEqual(len(list(checkpoint_dir.glob("*.error.json"))), 1)

            process_reports(
                cast(Any, FailingThenSuccessfulClient(should_fail=False)),
                [report_input],
                checkpoint_dir=checkpoint_dir,
            )

            self.assertEqual(len(list(checkpoint_dir.glob("*.error.json"))), 0)
            checkpoints = [path for path in checkpoint_dir.glob("*.json") if not path.name.endswith(".error.json")]
            self.assertEqual(len(checkpoints), 1)

    def test_process_reports_resume_skips_existing_checkpoint(self):
        class ResumeClient:
            def __init__(self, fail_on_alpha=False):
                self.fail_on_alpha = fail_on_alpha
                self.seen_texts = []

            def fetch_sdg_mentions(self, text):
                self.seen_texts.append(text)
                if self.fail_on_alpha and "alpha report" in text:
                    raise AssertionError("resume should not reprocess Alpha")
                return [{"sdg_number": 13, "mention_count": 1, "importance_score": 3.0, "page_numbers": [1]}]

            def categorize_investments(self, text, report_year=""):
                return {"keyword_categories": [], "investment_instruments": [], "unlinked_investments_eur": 0.0}

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            alpha = root / "Säästöpankki.pdf"
            beta = root / "Beta.pdf"
            (root / "Säästöpankki.pdf.parsed.json").write_text(
                json.dumps({"full_text": "alpha report", "pages": [[1, "alpha report"]]}), encoding="utf-8"
            )
            (root / "Beta.pdf.parsed.json").write_text(
                json.dumps({"full_text": "beta report", "pages": [[1, "beta report"]]}), encoding="utf-8"
            )
            checkpoint_dir = root / "checkpoints"
            alpha_report = ReportInput(str(alpha), bank="Säästöpankki", country="finland", year="2024")
            beta_report = ReportInput(str(beta), bank="Beta", country="finland", year="2024")
            process_reports(ResumeClient(), [alpha_report], checkpoint_dir=checkpoint_dir)

            client = ResumeClient(fail_on_alpha=True)
            mentions, investments = process_reports(
                client,
                [alpha_report, beta_report],
                checkpoint_dir=checkpoint_dir,
                resume=True,
            )

            self.assertIn("finland_Säästöpankki_2024", mentions)
            self.assertIn("finland_Beta_2024", investments)
            self.assertEqual(len(client.seen_texts), 1)
            self.assertIn("beta report", client.seen_texts[0])
            checkpoint_names = [path.name for path in checkpoint_dir.glob("*.json")]
            self.assertTrue(all("ää" not in name for name in checkpoint_names))


class DiscoveryTests(unittest.TestCase):
    def test_discover_pdfs_expands_directories_recursively(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            nested = root / "nested"
            nested.mkdir()
            first = root / "First.pdf"
            second = nested / "Second.PDF"
            ignored = root / "notes.txt"
            first.write_bytes(b"")
            second.write_bytes(b"")
            ignored.write_text("not a pdf", encoding="utf-8")

            discovered = discover_pdfs([str(root)])

            self.assertEqual(sorted(discovered), sorted([str(first), str(second)]))

    def test_manifest_requires_bank_value(self):
        with tempfile.TemporaryDirectory() as tmp:
            manifest = Path(tmp) / "reports.csv"
            manifest.write_text("path,bank\nreport.pdf,\n", encoding="utf-8")

            with self.assertRaisesRegex(RuntimeError, "empty bank"):
                load_report_manifest(manifest)


class MentionMergeTests(unittest.TestCase):
    def test_merge_mentions_matches_by_sdg_number_not_position(self):
        target = SdgMentions(
            [
                SdgMentionItem(sdg_number=1, mention_count=0, importance_score=0.0, page_numbers=[]),
                SdgMentionItem(sdg_number=7, mention_count=1, importance_score=1.0, page_numbers=[2]),
            ]
        )
        source = SdgMentions(
            [
                SdgMentionItem(sdg_number=7, mention_count=2, importance_score=4.0, page_numbers=[3, 2]),
                SdgMentionItem(sdg_number=1, mention_count=5, importance_score=2.0, page_numbers=[1]),
            ]
        )

        _merge_mentions(target, source)

        self.assertEqual(target.sdg_data[0].mention_count, 5)
        self.assertEqual(target.sdg_data[0].importance_score, 2.0)
        self.assertEqual(target.sdg_data[0].page_numbers, [1])
        self.assertEqual(target.sdg_data[1].mention_count, 3)
        self.assertEqual(target.sdg_data[1].importance_score, 4.0)
        self.assertEqual(target.sdg_data[1].page_numbers, [2, 3])


class ScoringTests(unittest.TestCase):
    def test_mismatch_uses_linked_sdg_amounts_only(self):
        mentions = SdgMentions(
            [
                SdgMentionItem(sdg_number=7, mention_count=1, importance_score=5.0, page_numbers=[2]),
            ]
        )
        investments = InvestmentCategorization(
            keyword_categories=[
                KeywordCategoryItem(
                    category="Renewable energy",
                    subcategories=[],
                    amount_eur=100.0,
                    relevant_sdgs=[7],
                    construct_type="disclosed_sdg_linked_finance",
                    sdg_support="explicit",
                    inclusion_policy="linked_scored",
                    score_eligible=True,
                )
            ],
            unlinked_investments_eur=900.0,
        )

        self.assertEqual(compute_mismatch(mentions, investments), 0.0)

    def test_mismatch_excludes_gross_context_rows_from_denominator(self):
        mentions = SdgMentions(
            [
                SdgMentionItem(sdg_number=7, mention_count=1, importance_score=5.0, page_numbers=[2]),
            ]
        )
        investments = InvestmentCategorization(
            keyword_categories=[
                KeywordCategoryItem(
                    category="Renewable energy",
                    subcategories=[],
                    amount_eur=100.0,
                    relevant_sdgs=[7],
                    construct_type="disclosed_sdg_linked_finance",
                    sdg_support="explicit",
                    inclusion_policy="linked_scored",
                    score_eligible=True,
                ),
                KeywordCategoryItem(
                    category="Facilitated finance",
                    subcategories=[],
                    amount_eur=900.0,
                    relevant_sdgs=[13],
                    construct_type="facilitated_sustainable_finance",
                    sdg_support="strong",
                    inclusion_policy="gross_context_only",
                    score_eligible=False,
                ),
            ],
        )

        self.assertEqual(compute_mismatch(mentions, investments), 0.0)


if __name__ == "__main__":
    unittest.main()
